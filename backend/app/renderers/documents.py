"""Document formats: PDF (reportlab), Word (python-docx), HTML, Markdown, text, JSON."""
from __future__ import annotations

import html
import io
import json
import re
from functools import lru_cache

from . import blocks as B
from .base import BRAND, FONT, FONT_BOLD, FONT_BOLD_ITALIC, FONT_ITALIC, FONT_MONO, RenderInput


def _footer_text(inp: RenderInput) -> str:
    m = inp.meta
    bits = ["BriefDesk", f"v{m['version']}" if m.get("version") else "", m.get("generated", "")]
    if m.get("verified_ratio") is not None:
        bits.append(f"{round(m['verified_ratio'] * 100)}% of citations verified")
    return " · ".join(b for b in bits if b)


# ---------------------------------------------------------------- PDF
@lru_cache(maxsize=1)
def _pdf_fonts() -> str:
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, path in [("BD", FONT), ("BD-Bold", FONT_BOLD), ("BD-Italic", FONT_ITALIC),
                       ("BD-BoldItalic", FONT_BOLD_ITALIC), ("BD-Mono", FONT_MONO)]:
        pdfmetrics.registerFont(TTFont(name, str(path)))
    addMapping("BD", 0, 0, "BD")
    addMapping("BD", 1, 0, "BD-Bold")
    addMapping("BD", 0, 1, "BD-Italic")
    addMapping("BD", 1, 1, "BD-BoldItalic")
    return "BD"


def _rl_markup(runs: list[B.Run]) -> str:
    out = []
    for text, st in runs:
        t = html.escape(text, quote=False)
        t = re.sub(r"\[(\d{1,3})\]", r'<font color="#2563eb" size="7"><super>[\1]</super></font>', t)
        if st.code:
            t = f'<font face="BD-Mono">{t}</font>'
        if st.bold:
            t = f"<b>{t}</b>"
        if st.italic:
            t = f"<i>{t}</i>"
        if st.link:
            t = f'<link href="{html.escape(st.link)}" color="#2563eb">{t}</link>'
        out.append(t)
    return "".join(out)


def render_pdf(inp: RenderInput) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        ListFlowable,
        ListItem,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font = _pdf_fonts()
    brand = colors.Color(*[c / 255 for c in BRAND])
    base = ParagraphStyle("base", fontName=font, fontSize=10.5, leading=15, alignment=TA_LEFT, spaceAfter=6)
    styles = {
        1: ParagraphStyle("h1", parent=base, fontName="BD-Bold", fontSize=21, leading=26, spaceAfter=4,
                          textColor=colors.HexColor("#0f172a")),
        2: ParagraphStyle("h2", parent=base, fontName="BD-Bold", fontSize=14, leading=18, spaceBefore=12,
                          spaceAfter=5, textColor=brand),
        3: ParagraphStyle("h3", parent=base, fontName="BD-Bold", fontSize=11.5, leading=15, spaceBefore=8),
    }
    meta_style = ParagraphStyle("meta", parent=base, fontSize=8.5, textColor=colors.HexColor("#64748b"))
    cell = ParagraphStyle("cell", parent=base, fontSize=9, leading=12, spaceAfter=0)
    quote = ParagraphStyle("quote", parent=base, leftIndent=10, textColor=colors.HexColor("#334155"),
                           fontName="BD-Italic")

    story: list = []
    first_h1 = True
    for blk in B.parse(inp.markdown):
        if isinstance(blk, B.Heading):
            story.append(Paragraph(_rl_markup(blk.runs), styles.get(blk.level, styles[3])))
            if blk.level == 1 and first_h1:
                first_h1 = False
                story.append(Paragraph(html.escape(_footer_text(inp)), meta_style))
                story.append(HRFlowable(width="100%", thickness=1.2, color=brand, spaceAfter=8))
        elif isinstance(blk, B.Paragraph):
            story.append(Paragraph(_rl_markup(blk.runs), base))
        elif isinstance(blk, B.Quote):
            story.append(Paragraph(_rl_markup(blk.runs), quote))
        elif isinstance(blk, B.ListBlock):
            items = [ListItem(Paragraph(_rl_markup(runs), base), leftIndent=12 + 12 * depth)
                     for depth, runs in blk.items]
            story.append(ListFlowable(items, bulletType="1" if blk.ordered else "bullet", start="1" if blk.ordered
                                      else "•", leftIndent=12, bulletFontName=font))
        elif isinstance(blk, B.TableBlock):
            data = ([[Paragraph(f"<b>{_rl_markup(c)}</b>", cell) for c in blk.header]] if blk.header else []) + \
                   [[Paragraph(_rl_markup(c), cell) for c in r] for r in blk.rows]
            if data:
                cols = max(len(r) for r in data)
                data = [r + [""] * (cols - len(r)) for r in data]
                t = Table(data, colWidths=[(A4[0] - 40 * mm) / cols] * cols, repeatRows=1 if blk.header else 0)
                t.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eff6ff") if blk.header else colors.white),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
                story += [t, Spacer(1, 6)]
        elif isinstance(blk, B.CodeBlock):
            story.append(Preformatted(blk.text, ParagraphStyle("code", parent=base, fontName="BD-Mono", fontSize=8.5)))
        elif isinstance(blk, B.Rule):
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cbd5e1")))

    footer = _footer_text(inp)

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#94a3b8"))
        canvas.drawString(20 * mm, 12 * mm, footer[:120])
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=20 * mm, title=inp.title, author="BriefDesk",
                            subject=inp.meta.get("topic", ""))
    doc.build(story or [Paragraph("(empty)", base)], onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# ---------------------------------------------------------------- Word
def _docx_runs(par, runs: list[B.Run]) -> None:
    from docx.shared import Pt, RGBColor

    for text, st in runs:
        for i, piece in enumerate(re.split(r"(\[\d{1,3}\])", text)):
            if not piece:
                continue
            r = par.add_run(piece)
            r.bold, r.italic = st.bold or None, st.italic or None
            if st.code:
                r.font.name = "Consolas"
            if i % 2 == 1:                              # citation marker
                r.font.superscript = True
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(*BRAND)
            elif st.link:
                r.font.color.rgb = RGBColor(*BRAND)
                r.underline = True


def render_docx(inp: RenderInput) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.core_properties.title, doc.core_properties.author = inp.title[:250], "BriefDesk"
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    for lvl in (1, 2, 3):
        doc.styles[f"Heading {lvl}"].font.color.rgb = RGBColor(*BRAND) if lvl > 1 else RGBColor(15, 23, 42)
    first_h1 = True
    for blk in B.parse(inp.markdown):
        if isinstance(blk, B.Heading):
            h = doc.add_heading(level=min(blk.level, 3) if blk.level > 1 else 0)
            _docx_runs(h, blk.runs)
            if blk.level == 1 and first_h1:
                first_h1 = False
                meta = doc.add_paragraph()
                r = meta.add_run(_footer_text(inp))
                r.font.size, r.font.color.rgb, r.italic = Pt(9), RGBColor(100, 116, 139), True
        elif isinstance(blk, (B.Paragraph, B.Quote)):
            p = doc.add_paragraph(style="Quote" if isinstance(blk, B.Quote) else None)
            _docx_runs(p, blk.runs)
        elif isinstance(blk, B.ListBlock):
            for depth, runs in blk.items:
                style = ("List Number" if blk.ordered else "List Bullet") + (" 2" if depth else "")
                _docx_runs(doc.add_paragraph(style=style), runs)
        elif isinstance(blk, B.TableBlock):
            cols = max([len(blk.header)] + [len(r) for r in blk.rows] or [1])
            table = doc.add_table(rows=0, cols=cols)
            table.style = "Light Grid Accent 1"
            for i, r in enumerate(([blk.header] if blk.header else []) + blk.rows):
                cells = table.add_row().cells
                for j in range(cols):
                    _docx_runs(cells[j].paragraphs[0], r[j] if j < len(r) else [])
                    if i == 0 and blk.header:
                        for run in cells[j].paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif isinstance(blk, B.CodeBlock):
            p = doc.add_paragraph()
            r = p.add_run(blk.text)
            r.font.name = "Consolas"
        elif isinstance(blk, B.Rule):
            doc.add_paragraph("―" * 30)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- HTML / Markdown / text / JSON
def _html_inline(runs: list[B.Run]) -> str:
    out = []
    for text, st in runs:
        t = html.escape(text)
        t = re.sub(r"\[(\d{1,3})\]", r'<sup class="cite"><a href="#src-\1">[\1]</a></sup>', t)
        if st.code:
            t = f"<code>{t}</code>"
        if st.bold:
            t = f"<strong>{t}</strong>"
        if st.italic:
            t = f"<em>{t}</em>"
        if st.link:
            t = f'<a href="{html.escape(st.link)}" rel="noopener noreferrer">{t}</a>'
        out.append(t)
    return "".join(out)


def render_html(inp: RenderInput) -> bytes:
    parts = []
    for blk in B.parse(inp.markdown):
        if isinstance(blk, B.Heading):
            parts.append(f"<h{blk.level}>{_html_inline(blk.runs)}</h{blk.level}>")
            if blk.level == 1:
                parts.append(f'<p class="meta">{html.escape(_footer_text(inp))}</p>')
        elif isinstance(blk, B.Paragraph):
            parts.append(f"<p>{_html_inline(blk.runs)}</p>")
        elif isinstance(blk, B.Quote):
            parts.append(f"<blockquote>{_html_inline(blk.runs)}</blockquote>")
        elif isinstance(blk, B.ListBlock):
            tag = "ol" if blk.ordered else "ul"
            parts.append(f"<{tag}>" + "".join(f"<li>{_html_inline(r)}</li>" for _, r in blk.items) + f"</{tag}>")
        elif isinstance(blk, B.TableBlock):
            head = "".join(f"<th>{_html_inline(c)}</th>" for c in blk.header)
            body = "".join("<tr>" + "".join(f"<td>{_html_inline(c)}</td>" for c in r) + "</tr>" for r in blk.rows)
            parts.append(f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")
        elif isinstance(blk, B.CodeBlock):
            parts.append(f"<pre><code>{html.escape(blk.text)}</code></pre>")
        elif isinstance(blk, B.Rule):
            parts.append("<hr>")
    body = "\n".join(parts)
    # anchor the numbered sources so [n] links jump to them
    body = re.sub(r"<li>(\d+)\. ", r'<li id="src-\1">\1. ', body)
    for s in inp.sources:
        href = html.escape(s["url"])
        body = body.replace(f'<a href="{href}"', f'<a id="src-{s["n"]}" href="{href}"', 1)
    doc = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; img-src data:">
<title>{html.escape(inp.title)}</title>
<style>
:root{{--ink:#0f172a;--muted:#64748b;--brand:#2563eb;--line:#e2e8f0;--bg:#fff}}
@media (prefers-color-scheme:dark){{:root{{--ink:#e2e8f0;--muted:#94a3b8;--brand:#60a5fa;--line:#334155;--bg:#0b1220}}}}
body{{margin:0;background:var(--bg);color:var(--ink);font:16px/1.65 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}}
main{{max-width:780px;margin:0 auto;padding:40px 20px 80px}}
h1{{font-size:2rem;line-height:1.2;margin:0 0 .3rem}} h2{{color:var(--brand);margin-top:2rem}}
.meta{{color:var(--muted);font-size:.85rem;border-bottom:2px solid var(--brand);padding-bottom:.8rem}}
a{{color:var(--brand)}} sup.cite a{{text-decoration:none;font-size:.75em}}
table{{border-collapse:collapse;width:100%;margin:1rem 0}} th,td{{border:1px solid var(--line);padding:.45rem .6rem;
text-align:left;vertical-align:top}} th{{background:color-mix(in srgb,var(--brand) 10%,transparent)}}
blockquote{{border-left:3px solid var(--brand);margin:1rem 0;padding:.2rem 1rem;color:var(--muted)}}
pre{{overflow:auto;background:color-mix(in srgb,var(--ink) 6%,transparent);padding:1rem;border-radius:8px}}
</style></head><body><main>
{body}
</main></body></html>"""
    return doc.encode("utf-8")


def render_md(inp: RenderInput) -> bytes:
    return (inp.markdown.rstrip() + f"\n\n---\n_{_footer_text(inp)}_\n").encode("utf-8")


def render_txt(inp: RenderInput) -> bytes:
    lines = []
    for blk in B.parse(inp.markdown):
        if isinstance(blk, B.Heading):
            text = B.plain(blk.runs)
            lines += ["", text.upper() if blk.level == 1 else text, ("=" if blk.level == 1 else "-") * len(text)]
        elif isinstance(blk, (B.Paragraph, B.Quote)):
            lines += [B.plain(blk.runs), ""]
        elif isinstance(blk, B.ListBlock):
            lines += [("  " * d) + ("- " if not blk.ordered else f"{i}. ") + B.plain(r)
                      for i, (d, r) in enumerate(blk.items, 1)] + [""]
        elif isinstance(blk, B.TableBlock):
            for r in ([blk.header] if blk.header else []) + blk.rows:
                lines.append(" | ".join(B.plain(c) for c in r))
            lines.append("")
        elif isinstance(blk, B.CodeBlock):
            lines += [blk.text, ""]
    for s in inp.sources:
        if f"{s['n']}. " not in "\n".join(lines):
            lines.append(f"[{s['n']}] {s['title']} {s['url']}")
    lines += ["", _footer_text(inp)]
    return "\n".join(lines).strip().encode("utf-8") + b"\n"


def render_json(inp: RenderInput) -> bytes:
    sections, current = [], None
    for blk in B.parse(inp.markdown):
        if isinstance(blk, B.Heading) and blk.level == 2:
            current = {"heading": B.plain(blk.runs), "content": []}
            sections.append(current)
        elif current is not None and not isinstance(blk, B.Heading):
            if isinstance(blk, B.ListBlock):
                current["content"] += [B.plain(r) for _, r in blk.items]
            elif isinstance(blk, B.TableBlock):
                current["content"].append({"table": [[B.plain(c) for c in r] for r in [blk.header] + blk.rows]})
            elif hasattr(blk, "runs"):
                current["content"].append(B.plain(blk.runs))
    data = {"title": inp.title, "meta": inp.meta, "sections": [s for s in sections if s["heading"] != "Sources"],
            "sources": inp.sources, "markdown": inp.markdown}
    return json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8")
